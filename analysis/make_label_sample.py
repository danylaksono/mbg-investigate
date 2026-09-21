"""Draw the sample for human labelling and write a BLIND labelling sheet plus a private key.

    python analysis/make_label_sample.py [--seed 20260920]

Purpose: measure how good the cause lexicon is, judged by people. The sheet a labeller sees carries the
sentence, its context and dropdown columns, and NOTHING about what the regexes decided (that would anchor
them). What the regexes decided, and which stratum a sentence came from, is in the key.

Design. Sentences come from articles linked to an incident (cited or independently found), from the
latest evidence run. They are drawn from strata so that BOTH questions can be answered:

  precision  Of sentences the rules call a finding, how many really are?   strata FLAG_*
  recall     Of real findings, how many do the rules miss?                 strata UNFLAGGED_*

  FLAG_CONTAMINATION   flagged as a lab finding of contamination
  FLAG_CLEAN           flagged as a lab result reported clean
  FLAG_AWAITING        flagged as awaiting a result
  UNFLAGGED_HIGH       NOT flagged, but names an agent (bacteria, nitrite ...) AND a lab word (hasil, uji,
                       sampel, BPOM ...): the small pool where missed findings are most likely
  UNFLAGGED_STRONG     NOT flagged, other strong cause vocabulary (penyebab, disebabkan, basi ...)
  UNFLAGGED_WEAK       NOT flagged, only weak cue words (hasil, periksa ...). A DIAGNOSTIC: its pool is huge, so
                       it is reported but never extrapolated into the recall estimate

Each stratum's pool size is stored in the key's frame file so estimates can be weighted. At most one
sentence per article and two per event are drawn, near-duplicate sentences (syndicated copy) are dropped,
and a few items are repeated under new ids to measure a labeller's own consistency.
"""

from __future__ import annotations

import argparse
import json
import random
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import nlp_explore as N  # noqa: E402
from recall_audit import STRONG  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "labelling"

AGENT_RE = re.compile(N.AGENT, re.I)
LAB_WORD = re.compile(r"hasil|\blab\b|laborat|labkesda|\buji\b|pengujian|sampel|bpom|periksa|pemeriksaan", re.I)
WEAK = re.compile(r"\bhasil\b|\blab\b|laborat|sampel|\buji\b|periksa|pemeriksaan|penyebab|racun", re.I)
QUOTAS = {"FLAG_CONTAMINATION": 18, "FLAG_CLEAN": 11, "FLAG_AWAITING": 5,
          "UNFLAGGED_HIGH": 24, "UNFLAGGED_STRONG": 14, "UNFLAGGED_WEAK": 8}
N_REPEATS = 4

LABELS = ["LAB_FINDING", "LAB_CLEAN", "AWAITING", "CAUSE_SPECULATION", "NONE"]
ABOUT = ["YES", "NO_OTHER_INCIDENT", "NO_GENERAL", "UNSURE"]


def norm_key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()[:120]


def stratum_of(sentence: str) -> tuple[str, str] | None:
    """(stratum, regex flag) or None when the sentence is not in the sampling frame."""
    if N.reports_lab_finding(sentence):
        return "FLAG_CONTAMINATION", "contamination"
    if N.reports_lab_clean(sentence):
        return "FLAG_CLEAN", "clean"
    if N.AWAITING.search(sentence):
        return "FLAG_AWAITING", "awaiting"
    if STRONG.search(sentence):
        return ("UNFLAGGED_HIGH" if AGENT_RE.search(sentence) and LAB_WORD.search(sentence) else "UNFLAGGED_STRONG"), "none"
    if WEAK.search(sentence):
        return "UNFLAGGED_WEAK", "none"
    return None


def load_frame(run: str):
    R = ROOT / "data" / "evidence" / "runs" / run
    decisions = {}
    for line in open(R / "link_decisions.jsonl", encoding="utf-8"):
        d = json.loads(line)
        if d["cited_events"] or d["decision"] == "linked":
            decisions[d["url"]] = d
    arts = {json.loads(l)["url"]: json.loads(l) for l in open(R / "articles.jsonl", encoding="utf-8")}
    ev = pd.read_csv(ROOT / "data/processed/events.csv").set_index("event_id")
    ev["venues"] = ev.venues.apply(json.loads)
    state = pd.read_csv(R / "event_summary.csv").set_index("event_id").state_all.to_dict()

    docs, inc, events, ledger = N.load()
    docs = N.attach_events(docs, inc, events)
    texts = {r["original_url"]: r["text"] for _, r in docs.iterrows()}
    for url, a in arts.items():
        if url not in texts and a.get("cache_path"):
            p = Path(a["cache_path"]).with_suffix("").with_suffix(".txt")
            if p.exists():
                texts[url] = N.RE_TEASER.sub("", p.read_text(encoding="utf-8"))
    return R, decisions, arts, ev, state, texts


def build_pools(decisions, arts, ev, state, texts):
    pools: dict[str, list[dict]] = {k: [] for k in QUOTAS}
    for url, d in decisions.items():
        text = texts.get(url)
        if not text:
            continue
        eid = (d["cited_events"] or d["event_ids"])[0]
        sents = N.sentences(text)
        for i, s in enumerate(sents):
            if not 40 < len(s) < 400:
                continue
            found = stratum_of(s)
            if not found:
                continue
            stratum, flag = found
            pools[stratum].append({"url": url, "event_id": eid, "i": i, "sentence": s, "flag": flag,
                                   "before": sents[i - 1][:320] if i else "", "after": sents[i + 1][:320] if i + 1 < len(sents) else ""})
    return pools


def draw(pools, seed):
    rng = random.Random(seed)
    chosen, used_url, used_key, per_event = [], set(), set(), {}
    for stratum, quota in QUOTAS.items():
        pool = pools[stratum][:]
        rng.shuffle(pool)
        n = 0
        for c in pool:
            key = norm_key(c["sentence"])
            if c["url"] in used_url or key in used_key or per_event.get(c["event_id"], 0) >= 2:
                continue
            used_url.add(c["url"]); used_key.add(key)
            per_event[c["event_id"]] = per_event.get(c["event_id"], 0) + 1
            chosen.append({**c, "stratum": stratum})
            n += 1
            if n == quota:
                break
    return chosen, rng


def describe_event(eid, ev):
    e = ev.loc[eid]
    venues = "; ".join(v for v in e.venues if v)[:110]
    return f"{e.kabkota_name}, {e.date_start} - {venues}" if venues else f"{e.kabkota_name}, {e.date_start}"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=20260920)
    args = ap.parse_args(argv)
    run = (ROOT / "data/evidence/LATEST").read_text().strip()
    R, decisions, arts, ev, state, texts = load_frame(run)
    pools = build_pools(decisions, arts, ev, state, texts)
    chosen, rng = draw(pools, args.seed)

    # repeat a few items under new ids: a labeller who disagrees with themselves is telling us something
    items = list(chosen) + rng.sample(chosen, N_REPEATS)
    rng.shuffle(items)
    sheet, key, first_seen = [], [], {}
    for n, it in enumerate(items, 1):
        item_id = f"L{n:03d}"
        a = arts[it["url"]]
        sheet.append({"item_id": item_id,
                      "event (place, date, venue)": describe_event(it["event_id"], ev),
                      "article title": a["title"][:150],
                      "outlet": urlparse(it["url"]).netloc.replace("www.", ""),
                      "article date": a.get("when", ""),
                      "sentence before": it["before"], "SENTENCE TO LABEL": it["sentence"], "sentence after": it["after"],
                      "url": it["url"], "label": "", "about this event?": "", "notes": ""})
        src = (it["url"], it["i"])
        key.append({"item_id": item_id, "stratum": it["stratum"], "regex_flag": it["flag"], "event_id": it["event_id"],
                    "event_state_all": state.get(it["event_id"]), "url": it["url"], "sentence_index": it["i"],
                    "repeat_of_item": first_seen.get(src)})            # the earlier-listed copy, if this is a repeat
        first_seen.setdefault(src, item_id)
    OUT.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(sheet).to_csv(OUT / "labelling_sheet.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(key).to_csv(OUT / "labelling_key_PRIVATE.csv", index=False, encoding="utf-8")
    frame = {"evidence_run": run, "seed": args.seed, "quotas": QUOTAS, "repeats": N_REPEATS,
             "pool_sizes": {k: len(v) for k, v in pools.items()}, "items": len(sheet),
             "drawn": pd.Series([k["stratum"] for k in key if not k["repeat_of_item"]]).value_counts().to_dict(),
             "labels": LABELS, "about": ABOUT}
    (OUT / "sampling_frame.json").write_text(json.dumps(frame, indent=2), encoding="utf-8")
    write_xlsx(pd.DataFrame(sheet), OUT / "labelling_sheet.xlsx")
    print(json.dumps(frame, indent=2))
    return 0


INSTRUCTIONS = """LABELLING INSTRUCTIONS

You are helping to measure how well an automatic rule finds statements about the CAUSE of school food-poisoning
incidents linked to Indonesia's Makan Bergizi Gratis (MBG) programme. Each row is one sentence from a news article,
with the sentence before and after it, and the incident the article was linked to.

Please label independently, without looking at any other labeller's sheet. Trust the sentence, not your guess about
what the article "probably" says. If you need more, open the url. About 30-40 minutes for the whole sheet.

COLUMN "label": choose ONE, the strongest that applies (top of the list wins)

  LAB_FINDING        The sentence states a LABORATORY or analytical RESULT showing contamination or a named agent
                     (bacteria such as E. coli / Salmonella / Bacillus, a chemical such as nitrite, histamine,
                     heavy metal) in food, water, vomit or swab samples.
                       "Hasil uji laboratorium menemukan bakteri E. coli pada nasi."
                       "Dari 14 sampel rectal swab, semuanya positif E-coli."
                     Not enough: an unnamed "dugaan" or a finding by smell/appearance ("ayam berbau basi").

  LAB_CLEAN          The sentence states a LABORATORY result reporting NO contamination or meeting standards.
                       "Hasil uji sampel air dinyatakan memenuhi syarat kesehatan."
                       "BPOM: tidak ditemukan bakteri E. coli pada sampel."

  AWAITING           Samples were taken or sent, or results are pending, or the cause is not yet known.
                       "Sampel makanan telah dikirim ke Labkesda."   "Masih menunggu hasil laboratorium."

  CAUSE_SPECULATION  A suspected cause or a general explanation, but NO lab result.
                       "Diduga karena ayamnya sudah basi."   "Keracunan biasanya disebabkan oleh bakteri."
                       Also a cause stated by an official WITHOUT mention of a lab test.

  NONE               Nothing about the cause or laboratory testing (symptoms, numbers, quotes, politics, ...).

COLUMN "about this event?"  Is the sentence about the incident in the "event" column?

  YES                 It concerns that incident (same place / school, same time).
  NO_OTHER_INCIDENT   It concerns a DIFFERENT incident (a comparison, a roundup, a "kasus lain").
  NO_GENERAL          A general statement not about any one incident.
  UNSURE              You cannot tell.

Answer both columns for every row. Use "notes" for anything odd. Some rows appear twice under different ids: label
each on its own, without going back. If a sentence is cut off or garbled, label what you can and say so in notes.

There are no trick questions and no "right" answer the sheet is hiding. Disagreements are useful. When you
finish, save the file and send it back unchanged in structure. Thank you.
"""


def write_xlsx(df: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    widths = {"item_id": 8, "event (place, date, venue)": 30, "article title": 34, "outlet": 16, "article date": 12,
              "sentence before": 40, "SENTENCE TO LABEL": 60, "sentence after": 40, "url": 26, "label": 20,
              "about this event?": 20, "notes": 30}
    for idx, col in enumerate(df.columns, 1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths.get(col, 20)
    head_fill = PatternFill("solid", fgColor="DDE7F0")
    target = PatternFill("solid", fgColor="FFF6D6")
    answer = PatternFill("solid", fgColor="E6F4E6")
    cols = {c: i for i, c in enumerate(df.columns, 1)}
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, len(df) + 2):
        for c in range(1, len(df.columns) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, cols["SENTENCE TO LABEL"]).fill = target
        ws.cell(r, cols["SENTENCE TO LABEL"]).font = Font(bold=True)
        ws.cell(r, cols["label"]).fill = answer
        ws.cell(r, cols["about this event?"]).fill = answer
    last = len(df) + 1
    for name, options in (("label", LABELS), ("about this event?", ABOUT)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Please choose one of the listed values.")
        ws.add_data_validation(dv)
        letter = ws.cell(1, cols[name]).column_letter
        dv.add(f"{letter}2:{letter}{last}")
    ws.freeze_panes = "B2"
    info = wb.create_sheet("Instructions", 0)
    for n, line in enumerate(INSTRUCTIONS.split("\n"), 1):
        info.cell(n, 1, line)
        info.cell(n, 1).alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 130
    info["A1"].font = Font(bold=True, size=13)
    wb.save(path)


if __name__ == "__main__":
    sys.exit(main())
